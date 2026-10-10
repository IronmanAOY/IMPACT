import os
import numpy as np
import pandas as pd
from typing import Optional, Union
from docx.shared import Inches
from docx import Document
import matplotlib.pyplot as plt

from impact_pipeline.analysis_bootstrap import (
    definedness_summary,
    holm_adjust,
    paired_session_test,
    paired_tests_table,
)

_SUPERSCRIPTS = str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹")
# ---------- formatting ----------


def fmt(x, decimals=3, sci_below=1e-4, sci_above=1e5):
    """Human-friendly float formatting. Uses scientific notation if |x| is
    very small/large to avoid '0.000' artifacts."""
    if x is None or (isinstance(x, float) and not np.isfinite(x)):
        return "na"
    ax = abs(x)
    if (ax != 0 and ax < sci_below) or ax >= sci_above:
        return f"{x:.3e}"
    return f"{x:.{decimals}f}"


def _save_paired_metric_plot(
    agg: pd.DataFrame,
    metric: str,
    out_path: str,
    title: str,
    y_label: str,
    scale: float = 1.0,
    sessions=("awake", "deep"),
):
    if metric not in agg.columns:
        return False

    piv = agg.pivot(index='subject', columns='session', values=metric)
    if not set(sessions).issubset(set(piv.columns)):
        return False

    paired = piv[[sessions[0], sessions[1]]].dropna()
    if paired.empty:
        return False

    a = paired[sessions[0]].to_numpy(dtype=float) * scale
    d = paired[sessions[1]].to_numpy(dtype=float) * scale
    means = [np.nanmean(a), np.nanmean(d)]
    sems = [
        np.nanstd(a, ddof=1) / np.sqrt(len(a)) if len(a) > 1 else np.nan,
        np.nanstd(d, ddof=1) / np.sqrt(len(d)) if len(d) > 1 else np.nan,
    ]

    x = np.array([0.0, 1.0], dtype=float)
    rng = np.random.RandomState(42)
    jitter = 0.045

    fig, ax = plt.subplots(figsize=(6.0, 4.0))
    for i in range(len(paired)):
        ax.plot(x, [a[i], d[i]], color="#B8B8B8", linewidth=0.7, alpha=0.6, zorder=1)

    ax.errorbar(
        x,
        means,
        yerr=sems,
        fmt='o',
        color="#1f4e79",
        ecolor="#1f4e79",
        elinewidth=1.5,
        capsize=4,
        markersize=7,
        zorder=3,
        label="Mean ± SEM",
    )

    ax.scatter(
        np.full_like(a, x[0]) + rng.uniform(-jitter, jitter, size=len(a)),
        a,
        color="#2a9d8f",
        alpha=0.75,
        s=24,
        zorder=2,
        label=sessions[0] if len(paired) > 0 else None,
    )
    ax.scatter(
        np.full_like(d, x[1]) + rng.uniform(-jitter, jitter, size=len(d)),
        d,
        color="#e76f51",
        alpha=0.75,
        s=24,
        zorder=2,
        label=sessions[1] if len(paired) > 0 else None,
    )

    ax.set_xticks(x)
    ax.set_xticklabels([sessions[0].title(), sessions[1].title()])
    ax.set_ylabel(y_label)
    ax.set_title(title)
    ax.grid(axis='y', linestyle=':', alpha=0.35)
    fig.tight_layout()
    fig.savefig(out_path, dpi=160)
    plt.close(fig)
    return True


def _theta_diff_table(
    df: pd.DataFrame, scale: float = 1.0, sessions=("awake", "deep")
) -> pd.DataFrame:
    """
    Build per-theta paired sessions[0] − sessions[1] S summary.
    """
    s0, s1 = str(sessions[0]), str(sessions[1])
    rows = []
    for theta, subdf in df.groupby('theta'):
        sub_mean = (
            subdf.assign(session=subdf['session'].astype(str))
            .groupby(['subject', 'session'], as_index=False)['S']
            .mean()
        )
        piv = sub_mean.pivot(index='subject', columns='session', values='S')
        if not {s0, s1}.issubset(set(piv.columns)):
            continue
        paired = piv[[s0, s1]].dropna()
        if paired.empty:
            continue
        a = paired[s0].to_numpy(dtype=float) * scale
        d = paired[s1].to_numpy(dtype=float) * scale
        diff = a - d
        n = int(np.isfinite(diff).sum())
        sem = float(np.nanstd(diff, ddof=1) / np.sqrt(n)) if n > 1 else np.nan
        rows.append(
            {
                'theta': float(theta),
                'mean_diff': float(np.nanmean(diff)),
                'sem_diff': sem,
                'n': n,
            }
        )
    return pd.DataFrame(rows)


def _iim_plot_spec(agg: pd.DataFrame):
    """(column, filename, title, y label, scale) for the IIM figure actually plotted."""
    # Use raw signed IIM (native units) for plots to preserve directionality/magnitude.
    raw_title = "Integrated Information Metric (IIM raw signed)"
    raw_label = "IIM raw (signed, unitless)"
    if 'IIM_raw' in agg.columns:
        return ("IIM_raw", "iim_curve.png", raw_title, raw_label, 1.0)
    if 'IIM_raw_scaled' in agg.columns:
        # Backward-compatible fallback for legacy tables.
        return (
            "IIM_raw_scaled", "iim_curve.png", raw_title, raw_label, 1.0
        )
    return (
        "IIM", "iim_curve.png", "Integrated Information Metric (IIM canonical)",
        "IIM canonical (0-1)", 1.0,
    )


def _save_threshold_difference_plot(
    df: pd.DataFrame,
    out_path: str,
    scale: float = 1.0,
    sessions=("awake", "deep"),
) -> bool:
    """
    Save the descriptive θ-scan plot of the paired sessions[0]−sessions[1] S
    difference.

    No θ is selected or marked: S is an exploratory legacy statistic and θ
    selection by the largest difference would bias any inference at that θ.
    Returns False (and writes nothing) when no paired θ summary is available.
    """
    ttab = _theta_diff_table(df, scale=scale, sessions=sessions)
    if ttab.empty:
        return False

    s0, s1 = str(sessions[0]), str(sessions[1])
    ttab = ttab.sort_values('theta')
    fig, ax = plt.subplots(figsize=(6.0, 4.0))
    ax.errorbar(ttab['theta'], ttab['mean_diff'], yerr=ttab['sem_diff'], marker='o')
    ax.axhline(0.0, color="#888888", linewidth=0.8)
    ax.set(xlabel='θ', ylabel=f'Mean S_{s0}–S_{s1} (exploratory)')
    fig.tight_layout()
    fig.savefig(out_path, dpi=160)
    plt.close(fig)
    return True


def _label_from_scale(name, scale):
    """Axis/label text such as 'S×10³' or 'CI×10⁻⁹' for pure powers of ten."""
    if scale == 1.0:
        return name
    try:
        exp = int(round(np.log10(scale)))
        if np.isclose(scale, 10.0 ** exp):
            sup = str(exp).translate(_SUPERSCRIPTS)
            return f"{name}×10{sup}"
    except Exception:
        pass
    return f"{name}×{scale:g}"


def _stat_line(label, st, scale=1.0, decimals=3, err="SEM"):
    """One paired-test sentence from a paired_session_test dict."""
    e_a = st.get('sem_a') if err == "SEM" else st.get('sd_a')
    e_b = st.get('sem_b') if err == "SEM" else st.get('sd_b')

    def sc(v):
        return v * scale if v is not None and np.isfinite(v) else np.nan

    def desc(ses, mean, e):
        return f"{ses} {fmt(sc(mean), decimals)} ± {fmt(sc(e), decimals)} ({err})"

    p_holm = st.get('p_holm')
    holm_txt = f", p_Holm={fmt(p_holm, decimals=4)}" if p_holm is not None else ""
    excl = int(st.get('n_subjects_excluded') or 0)
    excl_txt = (
        f" ({excl} subject(s) without a defined value in both sessions excluded)"
        if excl else ""
    )
    df_txt = st.get('df')
    df_txt = int(df_txt) if df_txt is not None and np.isfinite(df_txt) else "na"
    return (
        f"{label}: n_pairs={int(st.get('n_pairs') or 0)}{excl_txt}; "
        f"{desc(st.get('session_a', 'awake'), st.get('mean_a'), e_a)} vs "
        f"{desc(st.get('session_b', 'deep'), st.get('mean_b'), e_b)}; "
        f"t({df_txt})={fmt(st.get('t'))}, p={fmt(st.get('p'), decimals=4)}{holm_txt}, "
        f"dz={fmt(st.get('dz'))}."
    )


def _auc_line(name, s0, s1, res, extra=""):
    return (
        f"{name} AUC (P({name}_{s1} > {name}_{s0})) = {fmt(res.get('auc'))}, "
        f"subject-bootstrap 95% CI [{fmt(res.get('lo'))}, {fmt(res.get('hi'))}], "
        f"two-sided paired permutation p={fmt(res.get('p'), decimals=4)}{extra}."
    )


def _motion_paragraphs(motion):
    """Render motion_covariate_analysis output (one-row DataFrame or dict)."""
    if motion is None:
        return []
    if isinstance(motion, pd.DataFrame):
        if motion.empty:
            return ["Motion covariate analysis: no result."]
        motion = motion.iloc[0].to_dict()
    if not isinstance(motion, dict):
        kind = type(motion).__name__
        return [f"Motion covariate analysis: unsupported result type {kind}."]
    if 'skipped' in motion:
        return [f"Motion covariate analysis: skipped ({motion['skipped']})"]
    lines = []
    sessions = sorted(k[len('coef_'):] for k in motion if k.startswith('coef_'))
    for ses in sessions:
        coef = motion.get(f'coef_{ses}')
        pval = motion.get(f'p_{ses}')
        if isinstance(coef, (list, tuple)):
            coef = coef[0] if coef else np.nan
        if isinstance(pval, (list, tuple)):
            pval = pval[0] if pval else np.nan
        n_txt = f", n={motion.get(f'n_{ses}')}" if f'n_{ses}' in motion else ""
        lines.append(
            f"Motion (FD) within {ses}: coef={fmt(coef)}, "
            f"p={fmt(pval, decimals=4)}{n_txt}."
        )
    if 'delta_intercept' in motion:
        lines.append(
            "FD-adjusted state contrast "
            f"({motion.get('delta_sessions', 'awake-deep')}): "
            f"intercept={fmt(motion.get('delta_intercept'))}, "
            f"p={fmt(motion.get('p_delta_intercept'), decimals=4)}; "
            f"ΔFD coef={fmt(motion.get('delta_fd_coef'))}, "
            f"p={fmt(motion.get('p_delta_fd_coef'), decimals=4)}; "
            f"n_pairs={motion.get('n_pairs', 'na')}."
        )
    excl = []
    if motion.get('n_rows_undefined_score'):
        excl.append(f"{motion['n_rows_undefined_score']} row(s) with undefined score")
    if motion.get('n_subject_sessions_missing_fd'):
        n_fd = motion['n_subject_sessions_missing_fd']
        excl.append(f"{n_fd} subject-session(s) without mean_fd.txt")
    if excl:
        lines.append("Motion analysis exclusions: " + "; ".join(excl) + ".")
    return lines


def _ci_definedness_line(ci_ref, dfn):
    missing = dfn.get('missing_component_row_counts')
    missing_txt = f" (missing components: {missing})" if missing else ""
    # One row per run when definedness does not vary with θ (CI never does).
    unit = "runs" if dfn.get('count_unit') == "run" else "rows"
    return (
        f"CI reference: {ci_ref or 'unknown'} (each component divided by the "
        "reference mean; NAS enters directly, S is not part of CI). "
        f"{unit.capitalize()} with defined CI: {dfn.get('n_rows_defined', 'na')}/"
        f"{dfn.get('n_rows', 'na')}; "
        f"undefined {unit} excluded: {dfn.get('n_rows_undefined', 'na')}"
        f"{missing_txt}; "
        "subjects with defined CI in both sessions: "
        f"{dfn.get('n_subjects_complete_pairs', 'na')}."
    )


def _atlas_paragraphs(atlas_results, s0, s1):
    lines = []
    for name, r in atlas_results.items():
        is_dict = isinstance(r, dict)
        if is_dict and ("roughness" not in r) and ("skipped" in r or "reason" in r):
            reason = r.get('skipped', r.get('reason'))
            lines.append(f"Atlas {name}: skipped ({reason}).")
            continue
        if not (is_dict and ("roughness" in r)):
            try:
                lines.append(
                    f"Atlas {name}: metric roughness {s0}={fmt(r[s0], decimals=3)}, "
                    f"{s1}={fmt(r[s1], decimals=3)}."
                )
            except Exception:
                lines.append(f"Atlas {name}: {r}")
            continue
        rough = r.get("roughness") or {}
        if isinstance(rough, dict) and rough:
            items = ", ".join(f"{k}={fmt(v, decimals=3)}" for k, v in rough.items())
            lines.append(f"Atlas {name}: within-subject S roughness across θ {items}.")
        else:
            lines.append(f"Atlas {name}: S roughness not available.")
        metrics = r.get("metrics") if isinstance(r.get("metrics"), dict) else {}
        for m in ("S", "PDI", "NAS", "IIM", "IIM_raw", "CI", "RAM", "SRPI"):
            st = metrics.get(m)
            if not isinstance(st, dict) or int(st.get("n") or 0) <= 0:
                continue
            lines.append(
                f"Atlas {name} {m}: n={int(st.get('n'))}; "
                f"mean_{s0}={fmt(st.get('mean_a', np.nan), decimals=4)}, "
                f"mean_{s1}={fmt(st.get('mean_b', np.nan), decimals=4)}, "
                f"Δ({s0}−{s1})={fmt(st.get('delta_a_minus_b', np.nan), decimals=4)}, "
                f"t={fmt(st.get('t', np.nan))}, "
                f"p={fmt(st.get('p', np.nan), decimals=4)}, "
                f"p_Holm={fmt(st.get('p_holm', np.nan), decimals=4)}, "
                f"dz={fmt(st.get('d_paired', np.nan))}."
            )
        notes = r.get("notes")
        if isinstance(notes, list) and notes:
            lines.append(f"Atlas {name} notes: {' '.join(str(x) for x in notes)}")
    return lines


def _model_comparison_paragraphs(mc):
    lines = []
    for m, res in mc.items():
        try:
            ci_txt = ""
            ci = res.get('delta_auc_ci')
            if isinstance(ci, (list, tuple)) and len(ci) == 2:
                ci_txt = (
                    f", subject-bootstrap 95% CI [{fmt(ci[0], decimals=3)}, "
                    f"{fmt(ci[1], decimals=3)}]"
                )
            note = f" ({res['note']})" if res.get('note') else ""
            score = res.get('score', 'S')
            auc_txt = ""
            if 'auc_score' in res and 'auc_metric' in res:
                auc_txt = (
                    f"AUC_{score}={fmt(res['auc_score'], decimals=3)}, "
                    f"AUC_{m}={fmt(res['auc_metric'], decimals=3)}; "
                )
            disc_txt = ""
            if 'delta_discrimination' in res:
                dci = res.get('delta_discrimination_ci')
                dci_txt = ""
                if isinstance(dci, (list, tuple)) and len(dci) == 2:
                    dci_txt = (
                        f" [{fmt(dci[0], decimals=3)}, {fmt(dci[1], decimals=3)}]"
                    )
                disc_txt = (
                    " Discriminability Δ|AUC−0.5|="
                    f"{fmt(res['delta_discrimination'], decimals=3)}{dci_txt}, "
                    f"p={fmt(res.get('p_discrimination', np.nan), decimals=4)}, "
                    "p_Holm="
                    f"{fmt(res.get('p_discrimination_holm', np.nan), decimals=4)}."
                )
            lines.append(
                f"{score} vs {m}: {auc_txt}"
                f"signed ΔAUC={fmt(res['delta_auc'], decimals=3)}{ci_txt}, "
                f"p={fmt(res['p_val'], decimals=4)}, "
                f"p_Holm={fmt(res.get('p_holm', np.nan), decimals=4)}.{disc_txt}{note}"
            )
        except Exception:
            lines.append(f"Model {m}: {res}")
    return lines


# ---------- main API ----------

def create_doc(
    path: str,
    df: pd.DataFrame,
    df_stats_by_theta: pd.DataFrame,
    motion,
    atlas_results: dict,
    mc: dict,
    theta_results: Optional[dict] = None,
    fig_dir: Optional[str] = None,
    use_sem: bool = True,
    S_SCALE: float = 1e3,
    RAM_SCALE: float = 1e3,
    CI_SCALE: float = 1.0,      # scale applied to CI for display
    LABEL_S: Optional[str] = None,  # if None, auto from S_SCALE
    LABEL_RAM: Optional[str] = None,
    LABEL_CI: Optional[str] = None,
    repl: Optional[Union[dict, pd.DataFrame]] = None,
    stats: Optional[dict] = None,
    modality: Optional[str] = None,
    sessions=("awake", "deep"),
):
    """
    Build the Word report directly from the long 'df' produced by compute_synergy_ci
    and the other pipeline outputs.

    Parameters
    ----------
    df : DataFrame with columns ['subject','session','theta','S','CI','CI_defined',
        'CI_missing','CI_reference','RAM','PDI','NAS','IIM','SRPI'] (metrics repeat
        across theta; they are averaged per subject×session over defined rows only).
    df_stats_by_theta : DataFrame indexed by theta with ['t_S','p_S','d_S']
        (optionally 'mean_diff_S', 'p_S_holm').
    motion : DataFrame (one row) or dict returned by motion_covariate_analysis(...)
    atlas_results : dict from atlas_check(...)
    mc : dict from compare_models(...)
    theta_results : dict(theta -> {'auc_S','p_S'(,'p_S_holm')}), optional
    stats : dict from run_pipeline step 4 (CI definedness, bootstrap/permutation,
        Holm families)
    modality : 'fmri' / 'eeg' (report title)

    Inference rules: S (HypergraphSynergy) is an exploratory legacy statistic; it is
    summarised as the θ-grid average and per θ with Holm correction (no θ is selected
    post hoc). CI rows that are undefined are excluded (never treated as zero) and
    the exclusions are reported. Component tests are Holm-corrected as one family.
    """
    s0, s1 = str(sessions[0]), str(sessions[1])
    pair = (s0, s1)
    stats = dict(stats or {})
    err_label = "SEM" if use_sem else "SD"

    # ---------- derive per-subject/session means (defined rows only) ----------
    agg_map = dict(S=('S', 'mean'))
    for metric in ('CI', 'RAM', 'PDI', 'NAS', 'IIM', 'SRPI', 'IIM_raw', 'IIM_raw_scaled'):
        if metric in df.columns:
            agg_map[metric] = (metric, 'mean')
    agg = (df.groupby(['subject', 'session']).agg(**agg_map).reset_index())

    label_S = LABEL_S or _label_from_scale("S", S_SCALE)
    label_RAM = LABEL_RAM or _label_from_scale("RAM", RAM_SCALE)
    label_CI = LABEL_CI or _label_from_scale("CI", CI_SCALE)

    plot_specs = [
        ("S", "s_curve.png", "Synergy S (exploratory, θ-grid mean)", label_S, S_SCALE),
        ("CI", "ci_curve.png", "Consciousness Index (CI)", label_CI, CI_SCALE),
        ("RAM", "ram_curve.png", "Responsiveness-Adaptation Metric (RAM)",
         label_RAM, RAM_SCALE),
        ("PDI", "pdi_curve.png", "Phenomenal Differentiation Index (PDI)", "PDI", 1.0),
        ("NAS", "nas_curve.png", "Network Activation Synchrony (NAS)", "NAS", 1.0),
        ("SRPI", "srpi_curve.png", "Self-Referential Processing Index (SRPI)",
         "SRPI", 1.0),
    ]
    plot_specs.insert(5, _iim_plot_spec(agg))
    plot_specs = [spec for spec in plot_specs if spec[0] in agg.columns]

    # Generate per-metric session plots for the report when a figure directory
    # is provided. These are data-driven and use subject-level paired values.
    # Only figures drawn by this call are embedded, so a stale file from an earlier
    # run (e.g. the θ plot of run_s_ci, which marks the θ of the largest absolute
    # paired difference) is never reported.
    made_figs = set()
    if fig_dir:
        os.makedirs(fig_dir, exist_ok=True)
        theta_png = os.path.join(fig_dir, 'supp_theta_curve.png')
        if _save_threshold_difference_plot(df, theta_png, scale=1.0, sessions=pair):
            made_figs.add(theta_png)
        for metric, filename, title, ylabel, scale in plot_specs:
            fig_path = os.path.join(fig_dir, filename)
            if _save_paired_metric_plot(
                agg=agg,
                metric=metric,
                out_path=fig_path,
                title=title,
                y_label=ylabel,
                scale=scale,
                sessions=pair,
            ):
                made_figs.add(fig_path)

    # ---------- build document ----------
    doc = Document()
    modality_txt = None
    if modality:
        modality_txt = {"fmri": "fMRI", "eeg": "EEG"}.get(str(modality).lower())
    title = "Empirical Validation of MPC Metrics and CI"
    doc.add_heading(title + (f" ({modality_txt})" if modality_txt else ""), level=1)

    # 1) Primary: CI (three-valued; undefined rows excluded)
    if 'CI' in agg.columns:
        doc.add_heading("Consciousness Index (CI)", level=2)
        ci_ref = None
        if 'CI_reference' in df.columns and df['CI_reference'].notna().any():
            ci_ref = str(df['CI_reference'].dropna().iloc[0])
        dfn = stats.get('ci_definedness') or definedness_summary(df, sessions=pair)
        doc.add_paragraph(_ci_definedness_line(ci_ref, dfn))
        ci_test = stats.get('ci_test') or paired_session_test(agg, 'CI', pair)
        doc.add_paragraph(_stat_line(label_CI, ci_test, scale=CI_SCALE, err=err_label))
        auc_ci = stats.get('auc_CI')
        if isinstance(auc_ci, dict):
            n_used = auc_ci.get('n_subjects_used', 'na')
            extra = f" (n_subjects={n_used})"
            doc.add_paragraph(_auc_line("CI", s0, s1, auc_ci, extra))

    # 2) MPC components (one Holm family)
    comp_tab = stats.get('components')
    if isinstance(comp_tab, pd.DataFrame):
        comp_rows = comp_tab.to_dict('records')
    else:
        comps = [m for m in ('RAM', 'PDI', 'NAS', 'IIM', 'SRPI') if m in agg.columns]
        comp_rows = paired_tests_table(agg, comps, pair).to_dict('records')
    if comp_rows:
        doc.add_heading(
            "MPC components (paired, Holm-corrected across components)", level=2
        )
        comp_labels = {
            'RAM': (label_RAM, RAM_SCALE, 3),
            'PDI': ('PDI', 1.0, 4),
            'NAS': ('NAS', 1.0, 4),
            'IIM': ('IIM canonical (0-1)', 1.0, 4),
            'SRPI': ('SRPI', 1.0, 4),
        }
        for st in comp_rows:
            lab, sc, dec = comp_labels.get(st['metric'], (st['metric'], 1.0, 4))
            line = _stat_line(lab, st, scale=sc, decimals=dec, err=err_label)
            doc.add_paragraph(line)
    raw_cols = [c for c in ('IIM_raw', 'IIM_raw_scaled') if c in agg.columns]
    iim_raw_col = raw_cols[0] if raw_cols else None
    if iim_raw_col is not None:
        st_raw = paired_session_test(agg, iim_raw_col, pair)
        lab = "IIM raw (signed, unitless; supplementary, not in the Holm family)"
        doc.add_paragraph(_stat_line(lab, st_raw, decimals=4, err=err_label))

    # 3) Exploratory legacy statistic S
    doc.add_heading("Exploratory legacy statistic S (HypergraphSynergy)", level=2)
    s_test = stats.get('s_test') or paired_session_test(agg, 'S', pair)
    doc.add_paragraph(
        "S is not part of CI and is reported as an exploratory statistic. "
        "Headline summary: S averaged over the pre-declared θ grid; per-θ results "
        "are Holm-corrected across θ and no θ is selected post hoc."
    )
    doc.add_paragraph(
        _stat_line(f"{label_S} (θ-grid mean)", s_test, scale=S_SCALE, err=err_label)
    )
    auc_s = stats.get('auc_S')
    if isinstance(auc_s, dict):
        doc.add_paragraph(_auc_line("S", s0, s1, auc_s))
    if df_stats_by_theta is not None and not df_stats_by_theta.empty:
        doc.add_heading("S by θ (Holm across θ)", level=3)
        tab = df_stats_by_theta.sort_index()
        if 'p_S_holm' in tab.columns:
            p_holm = tab['p_S_holm'].to_numpy(dtype=float)
        elif 'p_S' in tab.columns:
            p_holm = holm_adjust(tab['p_S'].to_numpy(dtype=float))
        else:
            p_holm = np.full(len(tab), np.nan)
        for (theta, row), ph in zip(tab.iterrows(), p_holm):
            mean_diff_str = ""
            if "mean_diff_S" in tab.columns:
                mean_diff_str = f", mean_diff_S={fmt(getattr(row, 'mean_diff_S', np.nan))}"
            doc.add_paragraph(
                f"θ={fmt(theta, decimals=2)}: t_S={fmt(getattr(row, 't_S', np.nan))}, "
                f"p_S={fmt(getattr(row, 'p_S', np.nan), decimals=4)}, "
                f"p_Holm={fmt(ph, decimals=4)}, "
                f"dz_S={fmt(getattr(row, 'd_S', np.nan))}{mean_diff_str}"
            )

    # Permutation-test AUC of S by θ (two-sided, paired, Holm across θ)
    if isinstance(theta_results, dict) and len(theta_results):
        doc.add_heading(
            "Permutation-test AUC of S by θ (two-sided, within-subject, Holm across θ)",
            level=3,
        )
        thetas_sorted = sorted(theta_results)
        p_raw = [theta_results[t].get('p_S', np.nan) for t in thetas_sorted]
        p_adj = holm_adjust([np.nan if p is None else p for p in p_raw])
        for theta, ph in zip(thetas_sorted, p_adj):
            res = theta_results[theta]
            ph = res.get('p_S_holm', ph)
            doc.add_paragraph(
                f"θ={fmt(theta, decimals=2)}: "
                f"AUC_S={fmt(res.get('auc_S', np.nan), decimals=3)}, "
                f"p_S={fmt(res.get('p_S', np.nan), decimals=4)}, "
                f"p_Holm={fmt(ph, decimals=4)}"
            )

    # 4) Motion
    motion_lines = _motion_paragraphs(motion)
    if motion_lines:
        doc.add_heading("Motion covariates", level=2)
        for line in motion_lines:
            doc.add_paragraph(line)

    # 5) Atlas robustness
    if isinstance(atlas_results, dict) and atlas_results:
        doc.add_heading("Atlas robustness", level=2)
        for line in _atlas_paragraphs(atlas_results, s0, s1):
            doc.add_paragraph(line)

    # 6) Model comparisons
    if isinstance(mc, dict) and mc:
        doc.add_heading(
            "Comparison with baseline metrics (ΔAUC, paired DeLong, Holm)", level=2
        )
        for line in _model_comparison_paragraphs(mc):
            doc.add_paragraph(line)

    # 7) Independent replication
    if repl is not None:
        doc.add_heading("Independent Replication (exploratory)", level=2)

        def _write_rep_row(metric_name, stats_tuple):
            t, p, d_eff, dfree, mean_a, mean_d, err_a, err_d, n = stats_tuple
            doc.add_paragraph(
                f"{metric_name}: n={n}; "
                f"{s0} {fmt(mean_a, 4)}±{fmt(err_a, 4)} ({err_label}) vs "
                f"{s1} {fmt(mean_d, 4)}±{fmt(err_d, 4)} ({err_label}); "
                f"t({dfree})={fmt(t)}, p={fmt(p, decimals=4)}, dz={fmt(d_eff)}."
            )

        if isinstance(repl, dict):
            dataset = repl.get('dataset', 'independent dataset')
            doc.add_paragraph(f"Dataset: {dataset}")
            if {'delta_S', 'ci', 'cohend'}.issubset(repl.keys()):
                ci = repl['ci']
                ok = isinstance(ci, (list, tuple)) and len(ci) == 2
                lo, hi = ci if ok else (np.nan, np.nan)
                kind = repl.get('ci_kind', '95% CI')
                doc.add_paragraph(
                    f"ΔS={fmt(repl['delta_S'])} ({kind}: [{fmt(lo)}, {fmt(hi)}]), "
                    f"dz={fmt(repl['cohend'])}, n_pairs={repl.get('n_pairs', 'na')}."
                )
            else:
                doc.add_paragraph(
                    "(Replication summary keys not recognized — got: "
                    f"{list(repl.keys())})"
                )

        elif isinstance(repl, pd.DataFrame):
            agg_rep = {'S': ('S', 'mean')}
            if 'CI' in repl.columns:
                agg_rep['CI'] = ('CI', 'mean')
            rep = (repl.groupby(['subject', 'session'])
                       .agg(**agg_rep)
                       .reset_index())

            def _paired_block(metric):
                st = paired_session_test(rep, metric, pair)
                err_a = st['sem_a'] if use_sem else st['sd_a']
                err_d = st['sem_b'] if use_sem else st['sd_b']
                dfree = int(st['df']) if np.isfinite(st['df']) else 0
                return (st['t'], st['p'], st['dz'], dfree, st['mean_a'], st['mean_b'],
                        err_a, err_d, st['n_pairs'])

            _write_rep_row('Synergy (S)', _paired_block('S'))
            if 'CI' in rep.columns:
                _write_rep_row('Consciousness Index (CI)', _paired_block('CI'))

        else:
            doc.add_paragraph(f"(Unsupported replication object type: {type(repl)})")

    # 8) Figures
    if fig_dir:
        theta_fig = os.path.join(fig_dir, 'supp_theta_curve.png')
        if theta_fig in made_figs:
            doc.add_heading(
                "Supplementary: S difference across θ (exploratory, descriptive)",
                level=2,
            )
            doc.add_picture(theta_fig, width=Inches(5.0))
        else:
            doc.add_paragraph(
                "(θ-curve not drawn: no subjects with S in both "
                f"{s0} and {s1})"
            )

        headings = {"S": "Synergy S (exploratory)"}
        for metric, filename, title, _ylabel, _scale in plot_specs:
            fig_path = os.path.join(fig_dir, filename)
            if fig_path in made_figs:
                heading = f"Supplementary: {headings.get(metric, title)}"
                doc.add_heading(heading, level=2)
                doc.add_picture(fig_path, width=Inches(5.0))

    doc.save(path)
