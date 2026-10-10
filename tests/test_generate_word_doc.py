import numpy as np
import pandas as pd
import pytest
from docx import Document

from impact_pipeline import generate_word_doc as gwd


def _df(n=10, seed=0):
    rng = np.random.RandomState(seed)
    rows = []
    for i in range(n):
        base = rng.randn()
        for ses, shift in (("awake", 1.0), ("deep", 0.0)):
            ci = 1.0 + 0.3 * shift + rng.randn() * 0.1
            undefined = i == 0 and ses == "deep"
            for th in (0.3, 0.5, 0.7):
                rows.append(
                    {
                        "subject": f"s{i}",
                        "session": ses,
                        "theta": th,
                        "S": 1e-3 * (base + shift * th + rng.randn() * 0.1),
                        "CI": np.nan if undefined else ci,
                        "CI_defined": not undefined,
                        "CI_missing": "SRPI" if undefined else "",
                        "CI_reference": "cohort_high_state",
                        "RAM": 0.1 + 0.05 * shift,
                        "PDI": 0.02 + 0.01 * shift + rng.randn() * 0.001,
                        "NAS": 0.3 + rng.randn() * 0.01,
                        "IIM": 0.4,
                        "SRPI": np.nan if undefined else 0.5,
                        "IIM_raw": 0.2 + rng.randn() * 0.01,
                    }
                )
    return pd.DataFrame(rows)


def _text(path):
    return "\n".join(p.text for p in Document(str(path)).paragraphs)


def test_label_from_scale():
    # The exponent appears once, as a superscript (not 'S×10³3').
    assert gwd._label_from_scale("S", 1e3) == "S×10³"
    assert gwd._label_from_scale("CI", 1e9) == "CI×10⁹"
    assert gwd._label_from_scale("CI", 1e-9) == "CI×10⁻⁹"
    assert gwd._label_from_scale("X", 2.5) == "X×2.5"


def test_iim_plot_spec_matches_available_column():
    assert gwd._iim_plot_spec(pd.DataFrame(columns=["IIM"]))[2].endswith(
        "(IIM canonical)"
    )
    assert gwd._iim_plot_spec(pd.DataFrame(columns=["IIM", "IIM_raw"]))[0] == "IIM_raw"


def test_report_has_no_post_hoc_theta_and_reports_exclusions(tmp_path):
    df = _df()
    stats_by_theta = pd.DataFrame(
        {
            "t_S": [2.5, 3.0, 1.0],
            "p_S": [0.02, 0.01, 0.3],
            "d_S": [0.8, 0.9, 0.3],
            "mean_diff_S": [1, 2, 3],
        },
        index=pd.Index([0.3, 0.5, 0.7], name="theta"),
    )
    motion = pd.DataFrame(
        [
            {
                "coef_awake": 0.1,
                "p_awake": 0.5,
                "n_awake": 9,
                "coef_deep": -0.2,
                "p_deep": 0.4,
                "n_deep": 9,
                "delta_sessions": "awake-deep",
                "delta_intercept": 0.3,
                "p_delta_intercept": 0.001,
                "delta_fd_coef": 0.0,
                "p_delta_fd_coef": 0.9,
                "n_pairs": 9,
                "n_rows_undefined_score": 1,
                "n_subject_sessions_missing_fd": 0,
            }
        ]
    )
    atlas = {
        "aal90": {
            "roughness": {"awake": 0.1, "deep": 0.2},
            "metrics": {},
            "notes": ["n1"],
        },
        "shen268": {"skipped": "missing time series"},
    }
    mc = {
        "mean_conn": {
            "score": "S",
            "delta_auc": 0.1,
            "p_val": 0.04,
            "p_holm": 0.12,
            "delta_auc_ci": (-0.1, 0.3),
        }
    }
    theta_results = {
        0.3: {"auc_S": 0.3, "p_S": 0.01},
        0.5: {"auc_S": 0.2, "p_S": 0.02},
        0.7: {"auc_S": 0.4, "p_S": 0.5},
    }
    out = tmp_path / "report.docx"
    gwd.create_doc(
        path=str(out),
        df=df,
        df_stats_by_theta=stats_by_theta,
        motion=motion,
        atlas_results=atlas,
        mc=mc,
        theta_results=theta_results,
        fig_dir=str(tmp_path / "figs"),
        modality="eeg",
        repl={
            "dataset": "melb",
            "delta_S": 0.1,
            "ci": (0.05, 0.15),
            "cohend": 0.8,
            "n_pairs": 12,
            "ci_kind": "paired subject bootstrap 95% CI of mean difference",
        },
    )
    text = _text(out)
    assert "(EEG)" in text and "in fMRI" not in text
    assert "θ*" not in text and "largest |awake" not in text
    # Holm-adjusted p for S over θ: [0.02, 0.01, 0.3] -> [0.04, 0.03, 0.3]
    assert "θ=0.30: t_S=2.500, p_S=0.0200, p_Holm=0.0400" in text
    assert "p_Holm=0.0300" in text
    # CI exclusions are reported, not averaged as zeros; the one undefined run is
    # counted once (not once per θ).
    assert "Runs with defined CI: 19/20" in text
    assert "undefined runs excluded: 1" in text
    assert "1 subject(s) without a defined value in both sessions excluded" in text
    # The motion DataFrame is rendered.
    assert "FD-adjusted state contrast" in text and "Motion (FD) within awake" in text
    assert "Atlas shen268: skipped (missing time series)" in text
    assert "S vs mean_conn: signed ΔAUC=0.100" in text and "p_Holm=0.1200" in text
    assert "paired subject bootstrap 95% CI of mean difference" in text
    assert (tmp_path / "figs" / "supp_theta_curve.png").exists()


def test_component_family_is_holm_corrected_and_uses_complete_pairs(tmp_path):
    df = _df()
    out = tmp_path / "r.docx"
    gwd.create_doc(
        path=str(out),
        df=df,
        df_stats_by_theta=None,
        motion={"skipped": "x"},
        atlas_results={},
        mc={},
        fig_dir=None,
    )
    text = _text(out)
    srpi_line = next(line for line in text.split("\n") if line.startswith("SRPI:"))
    assert "n_pairs=9" in srpi_line and "p_Holm=" in srpi_line
    assert "Motion covariate analysis: skipped (x)" in text


def test_replication_dataframe_missing_session_does_not_crash(tmp_path):
    rep = pd.DataFrame(
        {"subject": ["a", "b"], "session": ["awake", "awake"], "S": [1.0, 2.0]}
    )
    out = tmp_path / "r.docx"
    gwd.create_doc(
        path=str(out),
        df=_df(),
        df_stats_by_theta=None,
        motion=None,
        atlas_results={},
        mc={},
        repl=rep,
    )
    assert "Synergy (S): n=0" in _text(out)


@pytest.mark.parametrize("scale", [1.0, 1e3])
def test_threshold_plot_has_no_selected_theta(tmp_path, scale):
    ok = gwd._save_threshold_difference_plot(
        _df(), str(tmp_path / "t.png"), scale=scale
    )
    assert ok is True and (tmp_path / "t.png").exists()


def _n_pictures(path):
    return len(Document(str(path)).inline_shapes)


def test_stale_figures_are_not_embedded(tmp_path):
    # A θ plot left by step 2 (which marks a post hoc θ*) must not be embedded when
    # create_doc cannot redraw it, e.g. for a session pair other than awake/deep.
    figs = tmp_path / "figs"
    figs.mkdir()
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots()
    ax.axvline(0.5, linestyle="--")
    fig.savefig(figs / "supp_theta_curve.png")
    fig.savefig(figs / "s_curve.png")
    plt.close(fig)
    df = _df().assign(session=lambda d: d["session"].map({"awake": "a", "deep": "b"}))
    out = tmp_path / "r.docx"
    gwd.create_doc(
        path=str(out), df=df, df_stats_by_theta=None, motion=None,
        atlas_results={}, mc={}, fig_dir=str(figs), sessions=("a", "light"),
    )
    text = _text(out)
    assert "θ-curve not drawn" in text
    assert _n_pictures(out) == 0
    # With the analysed pair available, the θ plot is redrawn and embedded.
    gwd.create_doc(
        path=str(out), df=df, df_stats_by_theta=None, motion=None,
        atlas_results={}, mc={}, fig_dir=str(figs), sessions=("a", "b"),
    )
    assert "S difference across θ" in _text(out) and _n_pictures(out) > 1


def test_model_comparison_line_reports_discriminability(tmp_path):
    mc = {
        "mean_conn": {
            "score": "S",
            "auc_score": 0.1,
            "auc_metric": 0.9,
            "delta_auc": -0.8,
            "p_val": 1e-6,
            "p_holm": 3e-6,
            "delta_auc_ci": (-0.9, -0.7),
            "delta_discrimination": 0.0,
            "delta_discrimination_ci": (-0.1, 0.1),
            "p_discrimination": 0.9,
            "p_discrimination_holm": 1.0,
        }
    }
    out = tmp_path / "r.docx"
    gwd.create_doc(
        path=str(out), df=_df(), df_stats_by_theta=None, motion=None,
        atlas_results={}, mc=mc,
    )
    line = next(t for t in _text(out).split("\n") if t.startswith("S vs mean_conn"))
    assert "AUC_S=0.100, AUC_mean_conn=0.900" in line
    assert "Discriminability Δ|AUC−0.5|=0.000 [-0.100, 0.100], p=0.9000" in line
